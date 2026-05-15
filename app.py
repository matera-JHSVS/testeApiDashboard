import streamlit as st
import pandas as pd
import os
import xml.etree.ElementTree as ET
import glob
import io
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# --- CONFIGURAÇÃO DE CAMINHO ---
DIRETORIO_RESULTS = r"C:\Desenvolvimento\robot\renda-fixa-testes-funcionais\robot-framework\acceptance-tests\api-renda-fixa\results"


@st.cache_data(ttl=5)
def carregar_dados_xml():
    lista_testes = []
    arquivos = glob.glob(os.path.join(DIRETORIO_RESULTS, "**/*.xml"), recursive=True)
    for arquivo in arquivos:
        if "sikuli" in arquivo.lower(): continue
        try:
            tree = ET.parse(arquivo)
            root = tree.getroot()
            nome_feature = os.path.basename(os.path.dirname(arquivo))
            for test in root.iter('test'):
                status = test.find('status').attrib.get('status', 'FAIL').upper()
                lista_testes.append({
                    "Feature": nome_feature,
                    "Cenário": test.attrib.get('name'),
                    "Status": status
                })
        except:
            continue
    return pd.DataFrame(lista_testes)


def gerar_excel_deploy(df):
    output = io.BytesIO()
    wb = Workbook()

    # Aba Resumo
    ws1 = wb.active
    ws1.title = "Resumo Geral"
    ws1.append(["Total Statistics (Feature)", "Total", "Pass", "Fail"])
    resumo = df.groupby('Feature')['Status'].value_counts().unstack(fill_value=0)
    for feat, row in resumo.iterrows():
        p, f = row.get('PASS', 0), row.get('FAIL', 0)
        ws1.append([feat, p + f, p, f])

    # Aba Falhas (Negrito nos CTs)
    ws2 = wb.create_sheet("Cenários com Falha")
    ws2.append(["Cenário com Falha", "Status"])
    df_fail = df[df['Status'] == 'FAIL'].sort_values('Feature')
    row_idx = 2
    for feat in df_fail['Feature'].unique():
        ws2.cell(row=row_idx, column=1, value=feat).font = Font(bold=True, italic=True)
        row_idx += 1
        for ct in df_fail[df_fail['Feature'] == feat]['Cenário']:
            cell = ws2.cell(row=row_idx, column=1, value=ct)
            cell.font = Font(bold=True)
            ws2.cell(row=row_idx, column=2, value="FAIL")
            row_idx += 1
    wb.save(output)
    return output.getvalue()


# --- INTERFACE ---
st.set_page_config(page_title="QA Dashboard", layout="wide")

# Cabeçalho limpo (apenas título, animação nativa do Streamlit já cuidará do resto)
st.title("🚀 Dashboard de Testes - Renda Fixa")

# Sidebar
with st.sidebar:
    st.title("⚙️ Painel")
    df = carregar_dados_xml()
    if not df.empty:
        if st.button("📊 Gerar Relatório", key="btn_exp"):
            excel = gerar_excel_deploy(df)
            st.download_button(
                label="📥 Baixar Planilha",
                data=excel,
                file_name=f"Relatorio_{datetime.now().strftime('%H%M')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="dl_btn"
            )

# Corpo do Dashboard
if not df.empty:
    st.divider()
    st.subheader("Resultados por Feature")

    # Aplicando width="stretch" para evitar logs de depreciação
    resumo_view = df.groupby('Feature')['Status'].value_counts().unstack(fill_value=0)
    st.dataframe(resumo_view, width="stretch")

else:
    st.warning("Aguardando dados...")

# Atualização automática
time.sleep(10)
st.rerun()