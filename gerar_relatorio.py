import os
import xml.etree.ElementTree as ET
import pandas as pd
import glob
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Configuração de Diretório
DIRETORIO_RESULTS = r"C:\Desenvolvimento\robot\renda-fixa-testes-funcionais\robot-framework\acceptance-tests\api-renda-fixa\results"
ARQUIVO_SAIDA = os.path.join(DIRETORIO_RESULTS, "relatorio_deploy_final.xlsx")


def gerar_relatorio_final():
    dados_sucesso = {}
    dados_falha = {}
    estatisticas = {}

    print(f"Varrendo {DIRETORIO_RESULTS}...")
    arquivos_xml = glob.glob(os.path.join(DIRETORIO_RESULTS, "**/*.xml"), recursive=True)

    for caminho_xml in arquivos_xml:
        if "sikuli" in caminho_xml.lower(): continue

        nome_feature = os.path.basename(os.path.dirname(caminho_xml))

        if nome_feature not in estatisticas:
            estatisticas[nome_feature] = {'Pass': 0, 'Fail': 0, 'Total': 0}
            dados_sucesso[nome_feature] = []
            dados_falha[nome_feature] = []

        try:
            tree = ET.parse(caminho_xml)
            root = tree.getroot()

            for test in root.iter('test'):
                nome_ct = test.attrib.get('name', 'Sem Nome')
                status_tag = test.find('status')

                if status_tag is not None:
                    res = status_tag.attrib.get('status', '').upper()
                    estatisticas[nome_feature]['Total'] += 1

                    if res == 'PASS':
                        estatisticas[nome_feature]['Pass'] += 1
                        dados_sucesso[nome_feature].append(nome_ct)
                    else:
                        estatisticas[nome_feature]['Fail'] += 1
                        dados_falha[nome_feature].append(nome_ct)
        except:
            continue

    # --- CRIAÇÃO DO EXCEL ---
    wb = Workbook()

    # Aba 1: Resumo Geral (Igual à sua imagem de 38 features)
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Geral"
    headers = ["Total Statistics (Feature)", "Total", "Pass", "Fail"]
    ws_resumo.append(headers)
    for cell in ws_resumo[1]: cell.font = Font(bold=True)

    for f in sorted(estatisticas.keys()):
        d = estatisticas[f]
        ws_resumo.append([f, d['Total'], d['Pass'], d['Fail']])

    # Linha ALL TESTS
    total_t = sum(d['Total'] for d in estatisticas.values())
    total_p = sum(d['Pass'] for d in estatisticas.values())
    total_f = sum(d['Fail'] for d in estatisticas.values())
    ws_resumo.append(["ALL TESTS (Resumo Geral)", total_t, total_p, total_f])
    for cell in ws_resumo[ws_resumo.max_row]: cell.font = Font(bold=True)

    # Aba 2: Cenários com Falha (Nomes dos CTs em NEGRITO)
    ws_falha = wb.create_sheet(title="Cenários com Falha")
    ws_falha.append(["Cenário com Falha", "Status"])
    ws_falha["A1"].font = Font(bold=True)
    ws_falha["B1"].font = Font(bold=True)

    row_f = 2
    fill_feat = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")

    for feature in sorted(dados_falha.keys()):
        if not dados_falha[feature]: continue

        # Linha da Feature (Separador)
        cell_feat = ws_falha.cell(row=row_f, column=1, value=feature)
        cell_feat.font = Font(bold=True, italic=True)  # Feature em negrito e itálico
        cell_feat.fill = fill_feat
        ws_falha.cell(row=row_f, column=2, value="FAIL").fill = fill_feat
        row_f += 1

        for ct in dados_falha[feature]:
            # NOME DO CT EM NEGRITO conforme solicitado
            cell_ct = ws_falha.cell(row=row_f, column=1, value=ct)
            cell_ct.font = Font(bold=True)
            ws_falha.cell(row=row_f, column=2, value="FAIL")
            row_f += 1

    # Aba 3: Cenários com Sucesso
    ws_sucesso = wb.create_sheet(title="Cenários com Sucesso")
    ws_sucesso.append(["Cenário", "Status"])
    row_s = 2
    for feature in sorted(dados_sucesso.keys()):
        if not dados_sucesso[feature]: continue
        ws_sucesso.cell(row=row_s, column=1, value=feature).font = Font(bold=True)
        row_s += 1
        for ct in dados_sucesso[feature]:
            ws_sucesso.cell(row=row_s, column=1, value=ct)
            ws_sucesso.cell(row=row_s, column=2, value="PASS")
            row_s += 1

    # Ajustes finais de largura
    for sheet in wb.worksheets:
        sheet.column_dimensions['A'].width = 90
        sheet.column_dimensions['B'].width = 15

    wb.save(ARQUIVO_SAIDA)
    print(f"✅ Relatório Final Gerado com {len(estatisticas)} features!")
    print(f"📍 Arquivo: {ARQUIVO_SAIDA}")


if __name__ == "__main__":
    gerar_relatorio_final()